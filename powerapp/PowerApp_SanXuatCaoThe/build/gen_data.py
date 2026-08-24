# -*- coding: utf-8 -*-
"""Xuat du lieu tab 'San xuat cao the' ra CSV + XLSX de nap vao Dataverse.

Tieu de cot trung KHOP TUYET DOI voi display name cua cot trong bang Dataverse,
nho vay buoc "Map columns" cua trinh import se tu dong khop.
"""
import csv, os
import openpyxl

HERE = os.path.dirname(os.path.abspath(__file__))
SRC = os.path.abspath(os.path.join(HERE, "..", "..", "Manufacturing.xlsx"))
SHEET = "Sản xuất cao thế"

HEADERS = ["Mã bản ghi", "Ngày sản xuất", "Thiết bị", "Cấp điện áp", "Dây chuyền",
           "SL kế hoạch", "SL thực tế", "SP lỗi", "SP đạt", "Giờ vận hành",
           "Giờ dừng máy", "Điện năng", "Giờ công", "Hoàn thành KH", "Tỷ lệ lỗi",
           "Năng suất", "Trạng thái", "Ghi chú"]


def rows():
    wb = openpyxl.load_workbook(SRC, data_only=True)
    ws = wb[SHEET]
    out = []
    for r in range(5, ws.max_row + 1):
        v = [ws.cell(r, c).value for c in range(1, 17)]
        if v[0] is None:
            continue
        (ngay, tb, kv, dc, kh, tt, loi, dat, vh, dm, kwh, gc, ht, tl, ns, tr) = v
        ma = "SX-%s-%s" % (ngay.strftime("%Y%m%d"), dc)
        out.append([
            ma, ngay.strftime("%Y-%m-%d"), tb, int(kv), dc,
            int(kh), int(tt), int(loi), int(dat),
            round(float(vh), 2), round(float(dm), 2), round(float(kwh), 2),
            round(float(gc), 2), round(float(ht), 4), round(float(tl), 4),
            round(float(ns), 4), tr, "",
        ])
    return out


def main():
    data = rows()
    csv_path = os.path.join(HERE, "SanXuatCaoThe_Data.csv")
    with open(csv_path, "w", encoding="utf-8-sig", newline="") as f:
        w = csv.writer(f)
        w.writerow(HEADERS)
        w.writerows(data)

    xlsx_path = os.path.join(HERE, "SanXuatCaoThe_Data.xlsx")
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "SanXuatCaoThe"
    ws.append(HEADERS)
    for row in data:
        ws.append(row)
    for c in range(2, len(HEADERS) + 1):
        pass
    for r in range(2, len(data) + 2):
        ws.cell(r, 2).number_format = "yyyy-mm-dd"
    wb.save(xlsx_path)

    print("rows =", len(data))
    print(csv_path)
    print(xlsx_path)


if __name__ == "__main__":
    main()
